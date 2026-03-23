from docx import Document

# 读取模板文件，检查占位符格式
template_file = "期初资料1/1 课程教学大纲基础模版.docx"
doc = Document(template_file)

print("检查模板文件中的占位符：")
print("-" * 50)

placeholders_found = []
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text
    if '{{' in text or '}}' in text:
        print(f"段落 {i+1}: {text}")
        placeholders_found.append(text)

print("\n所有找到的占位符：")
for placeholder in placeholders_found:
    print(placeholder)
